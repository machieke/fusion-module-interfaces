"""Compare an independent PowerPoint PDF export and audit the cited evidence.

This checks presentation artifacts only; it never reruns runtime integrations.
Run with the same optional requirements as build.py. The argument is a PDF
exported by LibreOffice (or another independent PowerPoint renderer).
"""
from __future__ import annotations

import argparse
from collections import Counter
import hashlib
import json
from pathlib import Path
import re
import subprocess
import unicodedata

import fitz
from PIL import Image, ImageDraw
from pptx import Presentation

import build


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def tokens(value: str) -> Counter:
    return Counter(re.findall(r"\S+", unicodedata.normalize("NFKC", value)))


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("independent_pdf", type=Path)
    args = parser.parse_args()
    root, deck, office = build.ROOT, build.HERE, args.independent_pdf.resolve().parent
    assert args.independent_pdf.resolve() != deck / "omega-asi-module-interfaces.pdf", "Use an independent export"
    source = fitz.open(deck / "omega-asi-module-interfaces.pdf")
    rendered = fitz.open(args.independent_pdf)
    prs = Presentation(deck / "omega-asi-module-interfaces.pptx")
    build.build_content()
    assert len(source) == len(rendered) == len(prs.slides) == len(build.SLIDES) == 44
    assert sum(not s.appendix for s in build.SLIDES) == build.MAIN_SLIDES == 36
    mismatches, bounds, previews, used_fonts = [], [], [], set()
    for i, (a, b, slide, model) in enumerate(zip(source, rendered, prs.slides, build.SLIDES), 1):
        if tokens(a.get_text()) != tokens(b.get_text()):
            mismatches.append(dict(slide=i, pdf_only=dict(tokens(a.get_text()) - tokens(b.get_text())),
                                   pptx_only=dict(tokens(b.get_text()) - tokens(a.get_text()))))
        assert tokens(a.get_text()) == tokens("\n".join(sh.text for sh in slide.shapes if sh.has_text_frame)), i
        assert slide.notes_slide.notes_text_frame.text == model.notes + "\n\nSource: " + model.source, i
        for label, page in [("direct", a), ("office", b)]:
            for block in page.get_text("dict")["blocks"]:
                for line in block.get("lines", []):
                    for span in line["spans"]:
                        used_fonts.add(span["font"])
                        r = fitz.Rect(span["bbox"])
                        if r.x0 < -.5 or r.y0 < -.5 or r.x1 > page.rect.width + .5 or r.y1 > page.rect.height + .5:
                            bounds.append(dict(slide=i, renderer=label, text=span["text"], bbox=span["bbox"]))
        path = office / f"slide-{i:02d}.png"
        b.get_pixmap(matrix=fitz.Matrix(1.2, 1.2), alpha=False).save(path)
        im = Image.open(path).convert("RGB")
        im.thumbnail((480, 270))
        previews.append(im)
    for offset in range(0, len(previews), 9):
        sheet = Image.new("RGB", (1488, 915), "#D5DCE0")
        draw = ImageDraw.Draw(sheet)
        for j, im in enumerate(previews[offset:offset+9]):
            x, y = 12+(j%3)*492, 22+(j//3)*300
            sheet.paste(im, (x, y))
            draw.text((x, y+274), f"PowerPoint {offset+j+1:02d}", fill="#102B3F")
        sheet.save(office / f"contact-{offset//9+1}.png")

    embedded = {}
    for page in source:
        for xref, _, _, name, *_ in page.get_fonts():
            if "DejaVu" in name:
                embedded[name] = bool(source.extract_font(xref)[3])
    assert len(embedded) == 3 and all(embedded.values())
    manifest = json.loads((deck / "build-manifest.json").read_text())
    for path, expected in manifest["evidence_sha256"].items():
        assert digest(root / path) == expected, path
        assert subprocess.check_output(["git", "show", f"{build.EVIDENCE_COMMIT}:{path}"], cwd=root) == (root / path).read_bytes(), path
    for path, expected in build.M31VERIFY["source"]["files"].items():
        assert digest(root / path) == expected, path
        assert hashlib.sha256(subprocess.check_output(
            ["git", "show", f"{build.BASE_COMMIT}:{path}"], cwd=root)).hexdigest() == expected, path
    # M3 measurements remain historical: audit against their own tested source.
    for path, expected in build.M3VERIFY["source"]["files"].items():
        assert hashlib.sha256(subprocess.check_output(
            ["git", "show", f"{build.M3RECORD['tested_source']['commit']}:{path}"], cwd=root)).hexdigest() == expected, path
    assert build.M3VERIFY["source"] == build.M3RECORD["tested_source"]
    assert build.M3VERIFY["source"]["status"] == []
    assert all(c["returncode"] == 0 for c in build.M3VERIFY["commands"])
    assert build.M3VERIFY["gates"] == build.M3RECORD["gates"]
    assert sum(v.startswith("PASS") for v in build.M3RECORD["gates"].values()) == 13
    assert build.M3RECORD["gates"]["CHAIN"] == "NOT_RUN"
    assert build.M3RECORD["blocked_prerequisites"] == []
    m2checksums = json.loads((build.M2 / "checksums.json").read_text())
    for directory, checksums in [(build.M2, m2checksums), (build.M3, build.M3RECORD["evidence_files"]),
                                 (build.M31, build.M31RECORD["archive_sha256"])]:
        for path, expected in checksums.items():
            assert digest(directory / path) == expected, path
    report = build.M31RECORD["report"]
    assert digest(build.M31 / report["path"]) == report["sha256"]
    assert build.M31VERIFY["source"]["commit"] == build.M31RECORD["tested_commit"] == build.BASE_COMMIT
    assert build.M31VERIFY["source"]["status"] == build.M31RECORD["tested_status"] == []
    assert all(c["returncode"] == 0 for c in build.M31VERIFY["commands"])
    assert build.M31VERIFY["gates"] == build.M31RECORD["gates"]
    assert sum(v.startswith("PASS") for v in build.M31RECORD["gates"].values()) == 12
    assert build.M31RECORD["gates"]["CHAIN"] == "NOT_RUN"
    assert build.M31RECORD["blocked_supported_prerequisites"] == []
    observed = build.M31VERIFY["observed"]
    assert observed["repeated_episode"] == dict(events=33, logical_observations=33,
        drain_batches=[16, 16, 1], pending_after_each_drain=0, backpressure_at=16)
    assert observed["feedback"]["native_calls_during_reconciliation"] == 0
    assert observed["feedback"]["original_result_and_budget"] == "unchanged"
    assert observed["feedback"]["attention_providers"] == ["sqlite", "atomic-file-fixture"]
    source_audit = json.loads((build.M31 / "source-audit.json").read_text())
    for path, expected in source_audit["unchanged"].items():
        assert digest(root / path) == expected, path
    probe = json.loads((build.M31 / "non-numeric-processing.json").read_text())
    assert probe["source"] == build.BASE_COMMIT and len(probe["cases"]) == 2
    for case in probe["cases"]:
        assert case["processing"] == "committed-and-acknowledged"
        assert case["numerical"] == "model-inapplicable" and case["truth"] is None
        assert case["observations"] == 1 and case["pending"] == 0
    closure_hive = next(c for c in build.M31VERIFY["commands"] if c["name"] == "closure-native-hive")
    assert closure_hive["tests"] == dict(tests=7, failures=0, errors=0, skipped=0)
    swap = json.loads((build.M12 / "m11-swap-demo.json").read_text())
    assert len(swap["runs"]) == 4
    for run in swap["runs"]:
        assert run["pending_outbox_at_end"] == run["acknowledged_data_loss"] == 0
        assert run["counts"] == dict(jobs=1, outbox=1, messages=2, atoms=2)
    assert len(build.COMPARISON["cases"]) == 4
    for case in build.COMPARISON["cases"]:
        assert case["measured_native_formula_invocations"] == 1
        assert case["successful_replay_new_native_invocations"] == 0
        assert case["selection_order"][0] == case["prefer"]
        assert case["terminal_reason"] == "rule-budget"
        assert case["counts"] == dict(completed=1, duplicate_admissions=0, examined=2,
                                      feedback_applied=1, feedback_duplicate=0, refused=0,
                                      reserved=1, unique_admissions=1)
    hive = next(c for c in build.M3VERIFY["commands"] if c["name"] == "m3-native-hive")
    assert hive["tests"] == dict(tests=8, failures=0, errors=0, skipped=0)
    result = dict(pages=len(source), speaker_notes_checked=len(prs.slides),
                  normalized_word_mismatches=mismatches, text_bounds_findings=bounds,
                  fonts=sorted(used_fonts), pdf_embedded_fonts=embedded,
                  evidence_inputs_unchanged=len(manifest["evidence_sha256"]),
                  archived_m2_files_checked=len(m2checksums),
                  archived_m3_files_checked=len(build.M3RECORD["evidence_files"]),
                  archived_m3_1_files_checked=len(build.M31RECORD["archive_sha256"]),
                  recorded_implementation_hashes_match=len(build.M31VERIFY["source"]["files"]),
                  historical_m3_implementation_hashes_checked=len(build.M3VERIFY["source"]["files"]),
                  preserved_boundary_files_checked=len(source_audit["unchanged"]),
                  gate_results=build.M31RECORD["gates"], independent_render_sha256=digest(args.independent_pdf))
    (office / "comparison.json").write_text(json.dumps(result, indent=2) + "\n")
    print(json.dumps(result, indent=2))
    assert not mismatches and not bounds


if __name__ == "__main__":
    main()
