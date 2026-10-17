"""Audit all module decks against source evidence and an independent PPTX export.

This validates presentation artifacts; it never runs integration or benchmark work.
"""
from __future__ import annotations
import argparse
from collections import Counter
import hashlib
import json
from pathlib import Path
import re
import subprocess
import unicodedata
import zipfile
import fitz
from PIL import Image, ImageDraw
from pptx import Presentation
import catalogue
import build_modules as build


def digest(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def tokens(text):return Counter(re.findall(r'\S+',unicodedata.normalize('NFKC',text)))


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('office_root',type=Path)
    args=parser.parse_args()
    root,here=build.ROOT,build.HERE
    series=json.loads((here/'series-manifest.json').read_text())
    assert series['decks']==15 and series['total_slides']==180
    for path,sha in series['input_sha256'].items():
        assert digest(root/path)==sha,path
        original=subprocess.check_output(['git','show',catalogue.SOURCE_COMMIT+':'+path],cwd=root)
        assert hashlib.sha256(original).hexdigest()==sha,path
    # Tie current module claims to the exact tested M5 source and archive.
    m5=json.loads((root/'evidence/m5/final/verification.json').read_text())
    assert m5['source']['commit']==catalogue.TESTED_COMMIT and m5['source']['status']==[]
    assert m5['result']=='PASS' and len(m5['gates'])==12 and set(m5['gates'].values())=={'PASS'}
    assert all(c['returncode']==0 for c in m5['commands'])
    for key in ['missing_prerequisites','executed_failures','observation_timeouts']:assert m5[key]==[]
    for path,sha in m5['source']['files'].items():assert digest(root/path)==sha,path
    archived=json.loads((root/'evidence/m5/checksums.json').read_text())
    for path,sha in archived.items():assert digest(root/'evidence/m5'/path)==sha,path
    # Assert each quantitative family cited in the benchmark slides.
    swap=json.loads((root/'docs/module-interfaces/evidence/m1_2/m11-swap-demo.json').read_text())
    assert len(swap['runs'])==4
    assert all(r['counts']==dict(jobs=1,outbox=1,messages=2,atoms=2) and r['pending_outbox_at_end']==0 for r in swap['runs'])
    demo=json.loads((root/'docs/module-interfaces/evidence/m2/final/demo.json').read_text())
    assert demo['native_revision']['strength']=='0.5' and demo['native_revision']['confidence']=='0.6666666666666666'
    assert demo['retained_alternatives']==2
    comp=json.loads((root/'docs/module-interfaces/evidence/m3/final/comparison.json').read_text())
    assert len(comp['cases'])==4
    assert [c['step_seconds'] for c in comp['cases']]==[.744361,.677157,.726126,.641150]
    assert all(c['measured_native_formula_invocations']==1 and c['successful_replay_new_native_invocations']==0 for c in comp['cases'])
    m31=json.loads((root/'docs/module-interfaces/evidence/m3_1/final/verification.json').read_text())
    assert m31['observed']['repeated_episode']==dict(events=33,logical_observations=33,drain_batches=[16,16,1],pending_after_each_drain=0,backpressure_at=16)
    assert m31['observed']['feedback']['native_calls_during_reconciliation']==0
    trace=json.loads((root/'evidence/m4/final/same-artifact.json').read_text())
    assert trace['bytes']==dict(input=53,result=143,signed_request=673)
    assert trace['final_observation']['evidence']['inclusion']['cost']==675
    lookup=json.loads((root/'evidence/m5/traces/lookup-trace-1.json').read_text())
    assert lookup['initial']==dict(payloads=8,query_pages=3,relations=15,roots=5) and lookup['catchup']['roots']==6
    results=[]
    for m in catalogue.MODULES:
        folder=here/m['id'];office=args.office_root/m['id'];office.mkdir(parents=True,exist_ok=True)
        source=fitz.open(folder/'slides.pdf');other=fitz.open(office/'slides.pdf');ppt=Presentation(folder/'slides.pptx')
        model=build.build_module(m)
        assert len(source)==len(other)==len(ppt.slides)==len(model)==12,m['id']
        mismatches=[];bounds=[];previews=[];fonts={}
        for i,(a,b,slide,spec) in enumerate(zip(source,other,ppt.slides,model),1):
            assert tokens(a.get_text())==tokens('\n'.join(sh.text for sh in slide.shapes if sh.has_text_frame)),(m['id'],i)
            if tokens(a.get_text())!=tokens(b.get_text()):mismatches.append(i)
            assert slide.notes_slide.notes_text_frame.text==spec.notes+'\n\nSource: '+spec.source,(m['id'],i)
            for label,page in [('direct',a),('office',b)]:
                for block in page.get_text('dict')['blocks']:
                    for line in block.get('lines',[]):
                        for span in line['spans']:
                            r=fitz.Rect(span['bbox'])
                            if r.x0<-.5 or r.y0<-.5 or r.x1>page.rect.width+.5 or r.y1>page.rect.height+.5:bounds.append([i,label,span['text']])
            for xref,_,_,name,*_ in a.get_fonts():
                if 'DejaVu' in name:fonts[name]=bool(source.extract_font(xref)[3])
            path=office/f'slide-{i:02d}.png';b.get_pixmap(matrix=fitz.Matrix(1.2,1.2),alpha=False).save(path)
            im=Image.open(path).convert('RGB');im.thumbnail((480,270));previews.append(im)
        assert all(fonts.values()) and len(fonts)==3
        for offset in range(0,12,9):
            sheet=Image.new('RGB',(1488,915),'#D5DCE0');d=ImageDraw.Draw(sheet)
            for j,im in enumerate(previews[offset:offset+9]):
                x=12+(j%3)*492;y=22+(j//3)*300;sheet.paste(im,(x,y));d.text((x,y+274),f'{m["id"]} / {offset+j+1:02d}',fill='#102B3F')
            sheet.save(office/f'contact-{offset//9+1}.png')
        notes=(folder/'speaker-notes.md').read_text()
        for target in re.findall(r'\]\(([^)]+)\)',notes):
            if target.startswith('https://github.com/machieke/fusion-module-interfaces/blob/'):
                prefix='https://github.com/machieke/fusion-module-interfaces/blob/'+catalogue.SOURCE_COMMIT+'/'
                assert target.startswith(prefix) and (root/target[len(prefix):]).exists(),target
            else:assert (folder/target).exists(),target
        result=dict(module=m['id'],pages=12,speaker_notes=12,text_mismatches=mismatches,bounds_findings=bounds,embedded_fonts=fonts,
                    sha256={n:digest(folder/n) for n in ['slides.pdf','slides.pptx','speaker-notes.md','build-manifest.json']},
                    independent_pdf_sha256=digest(office/'slides.pdf'))
        results.append(result);assert not mismatches and not bounds,result
        print(m['id'],'PASS',flush=True)
    combined=fitz.open(here/'all-modules.pdf')
    assert len(combined)==180 and len([r for r in combined.get_toc() if r[0]==1])==15
    with zipfile.ZipFile(here/'module-decks.zip') as archive:
        assert archive.testzip() is None
        for name in ['catalogue.py','build_modules.py','render.py','check_modules.py','series-manifest.json']:
            assert archive.read('source/'+name)==(here/name).read_bytes(),name
        assert archive.read('fonts/LICENSE.txt')==(here.parent/'fonts/LICENSE.txt').read_bytes()
        for target in re.findall(r'\]\(([^)]+)\)',archive.read('README.md').decode()):
            if target.endswith('/'):assert any(p.startswith(target) for p in archive.namelist()),target
            else:assert target in archive.namelist(),target
        for m in catalogue.MODULES:
            for name in ['slides.pdf','slides.pptx','speaker-notes.md']:
                assert archive.read(m['id']+'/'+name)==(here/m['id']/name).read_bytes()
    result=dict(schema='fusion-module-talks-qa/1',date='2026-10-17',decks=15,pages=180,
        evidence_inputs_checked=len(series['input_sha256']),current_runtime_hashes_checked=len(m5['source']['files']),
        m5_archive_files_checked=len(archived),m5_gates=m5['gates'],module_results=results,
        combined_pdf_pages=180,combined_pdf_module_bookmarks=15,bundle_checked=True,
        runtime_or_benchmark_tests_executed_for_decks=False,
        renderer_limit='Independent LibreOffice and direct PDF checked; Microsoft PowerPoint itself unavailable.',
        files_sha256={p.name:digest(p) for p in [here/'catalogue.py',here/'build_modules.py',here/'render.py',here/'check_modules.py',here/'all-modules.pdf',here/'module-decks.zip',here/'series-manifest.json']})
    (here/'qa.json').write_text(json.dumps(result,indent=2)+'\n')
    print('PASS: 15 decks, 180 slides, independent rendering, notes, sources and bundle.')

if __name__=='__main__':main()
